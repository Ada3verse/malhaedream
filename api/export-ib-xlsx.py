import json
import pathlib
import re
from datetime import datetime
from http.server import BaseHTTPRequestHandler
from io import BytesIO
from urllib.parse import quote

from openpyxl import load_workbook

TEMPLATE_PATH = str(pathlib.Path(__file__).resolve().parent / 'ib_unit_plan_template.xlsx')
SHEET_NAME = 'UNIT 4.'

GRADE_MAP = {
    'MYP 1': '1학년',
    'MYP 2': '2학년',
    'MYP 3': '3학년',
}


def myp_year_to_grade(myp_year):
    if not myp_year:
        return '1학년'
    if myp_year in GRADE_MAP:
        return GRADE_MAP[myp_year]
    for digit, grade in (('1', '1학년'), ('2', '2학년'), ('3', '3학년')):
        if digit in str(myp_year):
            return grade
    return '1학년'


def set_cell(ws, ref, value):
    if value is None or value == '':
        return
    ws[ref] = str(value)


def extract_inquiry_questions(text):
    if not text:
        return {}
    factual = re.search(r'사실적[^\n]*\n([\s\S]+?)(?=개념적|논쟁적|$)', text)
    conceptual = re.search(r'개념적[^\n]*\n([\s\S]+?)(?=사실적|논쟁적|$)', text)
    debatable = re.search(r'논쟁적[^\n]*\n([\s\S]+?)(?=사실적|개념적|$)', text)
    return {
        'factual': factual.group(1).strip() if factual else '',
        'conceptual': conceptual.group(1).strip() if conceptual else '',
        'debatable': debatable.group(1).strip() if debatable else '',
    }


def extract_grasps(text):
    if not text:
        return {}
    patterns = {
        'goal': r'Goal\s*[\(（]?목표[\)）]?\s*[:：]\s*([^\n]+(?:\n(?!Role|Audience|Situation|Perform|Standard)[^\n]+)*)',
        'role': r'Role\s*[\(（]?역할[\)）]?\s*[:：]\s*([^\n]+(?:\n(?!Goal|Audience|Situation|Perform|Standard)[^\n]+)*)',
        'audience': r'Audience\s*[\(（]?청중[\)）]?\s*[:：]\s*([^\n]+(?:\n(?!Goal|Role|Situation|Perform|Standard)[^\n]+)*)',
        'situation': r'Situation\s*[\(（]?상황[\)）]?\s*[:：]\s*([^\n]+(?:\n(?!Goal|Role|Audience|Perform|Standard)[^\n]+)*)',
        'performance': r'Perform(?:ance|ace)?\s*[\(（]?수행[\)）]?\s*[:：]\s*([^\n]+(?:\n(?!Goal|Role|Audience|Situation|Standard)[^\n]+)*)',
        'standards': r'Standard[s]?\s*[\(（]?평가[\)）]?\s*[:：]\s*([^\n]+(?:\n(?!Goal|Role|Audience|Situation|Perform)[^\n]+)*)',
    }
    result = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            result[key] = match.group(1).strip()
    return result


ATL_CATEGORY_KEYWORDS = [
    '의사소통', '대인 관계 기능', '자기 관리 기능', '조사 기능', '사고 기능',
    'Communication', 'Social', 'Self-management', 'Research', 'Thinking',
]


def parse_atl_items(text):
    if not text:
        return []
    results = []
    current = None

    for raw_line in text.split('\n'):
        line = raw_line.strip()
        if not line or line.startswith('#'):
            continue

        is_category = any(k in line for k in ATL_CATEGORY_KEYWORDS) and len(line) < 60

        if is_category:
            if current:
                results.append(current)
            current = {'category': line, 'cluster': '', 'detail': ''}
        elif current is not None:
            if not current['cluster'] and len(line) < 80:
                current['cluster'] = line
            else:
                current['detail'] = (current['detail'] + '\n' + line) if current['detail'] else line

    if current:
        results.append(current)

    if not results and text.strip():
        results.append({'category': '', 'cluster': '', 'detail': text.strip()})

    return results


def build_workbook(info, sections):
    wb = load_workbook(TEMPLATE_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f'시트 "{SHEET_NAME}"를 찾을 수 없습니다.')
    ws = wb[SHEET_NAME]

    statement_prompt = (sections.get('statement') or {}).get('prompt') or ''
    inquiry_prompt = (sections.get('inquiry_questions') or {}).get('prompt') or ''
    assessment_prompt = (sections.get('assessment') or {}).get('prompt') or ''
    atl_prompt = (sections.get('atl') or {}).get('prompt') or ''
    formative_prompt = (sections.get('formative') or {}).get('prompt') or ''

    set_cell(ws, 'B3', info.get('teacherName') or '')
    set_cell(ws, 'E3', info.get('subject') or '')
    set_cell(ws, 'B4', info.get('unitTitle') or '')
    set_cell(ws, 'E4', myp_year_to_grade(info.get('mypYear')))

    related_concepts = info.get('relatedConcepts')
    if isinstance(related_concepts, list):
        related_concepts = ', '.join(related_concepts)

    set_cell(ws, 'A9', info.get('keyConcept') or '')
    set_cell(ws, 'C9', related_concepts or '')
    set_cell(ws, 'E9', info.get('globalContext') or '')
    set_cell(ws, 'F10', info.get('contextExploration') or '')

    set_cell(ws, 'A12', statement_prompt)

    if inquiry_prompt:
        questions = extract_inquiry_questions(inquiry_prompt)
        set_cell(ws, 'A15', questions.get('factual'))
        set_cell(ws, 'A17', questions.get('conceptual'))
        set_cell(ws, 'A19', questions.get('debatable'))

    if assessment_prompt:
        grasps = extract_grasps(assessment_prompt)
        set_cell(ws, 'D22', grasps.get('goal'))
        set_cell(ws, 'D23', grasps.get('role'))
        set_cell(ws, 'D24', grasps.get('audience'))
        set_cell(ws, 'D25', grasps.get('situation'))
        set_cell(ws, 'D26', grasps.get('performance'))
        set_cell(ws, 'D27', grasps.get('standards'))

    if atl_prompt:
        atl_items = parse_atl_items(atl_prompt)
        atl_rows = [30, 33]
        for row, item in zip(atl_rows, atl_items):
            set_cell(ws, f'A{row}', item.get('category'))
            set_cell(ws, f'B{row}', item.get('cluster'))
            set_cell(ws, f'C{row}', item.get('detail'))

    set_cell(ws, 'E42', formative_prompt)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


class handler(BaseHTTPRequestHandler):
    def _set_cors_headers(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')

    def do_OPTIONS(self):
        self.send_response(204)
        self._set_cors_headers()
        self.end_headers()

    def do_POST(self):
        try:
            content_length = int(self.headers.get('Content-Length', 0))
            raw_body = self.rfile.read(content_length) if content_length else b'{}'
            payload = json.loads(raw_body or b'{}')
            info = payload.get('info') or {}
            sections = payload.get('sections') or {}

            xlsx_bytes = build_workbook(info, sections)

            subject = (info.get('subject') or '교과').strip() or '교과'
            title = (info.get('unitTitle') or '단원').strip() or '단원'
            date_str = datetime.now().strftime('%Y%m%d')
            filename = f'IB_단원계획서_{subject}_{title}_{date_str}.xlsx'
            filename = re.sub(r'[/\\?%*:|"<>]', '_', filename)
            encoded_filename = quote(filename)

            self.send_response(200)
            self._set_cors_headers()
            self.send_header(
                'Content-Type',
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            self.send_header(
                'Content-Disposition',
                f"attachment; filename=\"ib_unit_plan.xlsx\"; filename*=UTF-8''{encoded_filename}",
            )
            self.send_header('Content-Length', str(len(xlsx_bytes)))
            self.end_headers()
            self.wfile.write(xlsx_bytes)
        except Exception as exc:  # noqa: BLE001
            error_body = json.dumps({'error': str(exc)}).encode('utf-8')
            self.send_response(500)
            self._set_cors_headers()
            self.send_header('Content-Type', 'application/json')
            self.send_header('Content-Length', str(len(error_body)))
            self.end_headers()
            self.wfile.write(error_body)
